import io
from datetime import date

import openpyxl
import pandas as pd
import pytest

import ingest


def xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Layouts copied from the week 38/2026 exports.
HAPANA_MEMBERS = xlsx([
    ["Export info"],
    ["Exported at", "21/09/2026, 5:37 am (GMT+00:00)"],
    ["Filters"],
    ["Date range", "14/09/2026 - 20/09/2026"],
    ["Weekly Membership"],
    ["Success Tutoring - Business name", "Date - Week/Year", "# Active members",
     "# Suspended members", "# Cancelled members", "# New members"],
    ["Success Tutoring - Auburn", "38/2026", 83, 9, 3, 1],
    ["Success Tutoring - Epping", "38/2026", 100, 12, 0, 2],
    ["Success Tutoring - Epping VIC", "38/2026", 46, 10, 0, 1],
    ["Success Tutoring - Belmont", "38/2026", 17, 3, 0, 0],
    ["Success Tutoring - Success Tutoring Sandbox", "38/2026", 35, 0, 0, 0],
    ["Success Tutoring - Howick", "38/2026", 19, 3, 0, 0],
])
HAPANA_REV_AU = xlsx([
    ["SA - Revenue Location Summary"],
    ["Success Tutoring - Business name", "Date - Week/Year",
     "# Active members - Date range end", "Total Sessions",
     "Average Gross Revenue per Location"],
    ["Success Tutoring - Auburn", "38/2026", 83, 36, "5,970"],
    ["Success Tutoring - Epping", "38/2026", 100, 23, "7,152"],
    ["Success Tutoring - Epping VIC", "38/2026", 46, 22, "3,533"],
    ["Success Tutoring - Belmont", "38/2026", 17, 13, "1,088"],
])
HAPANA_REV_NZ = xlsx([
    ["SA - Revenue Location Summary"],
    ["Business name", "Date - Week/Year", "# Active members - Date range end",
     "Total Sessions", "Average Gross Revenue per Location"],
    ["Howick", "38/2026", 19, 17, "1,346"],
])
INHOUSE_MEMBERS = xlsx([
    ["Members"],
    ["Location Name", "Active Members", "Suspended Members", "Cancelled Members",
     "New Members"],
    ["Success Tutoring - Burwood", 56, 11, 3, 1],
    ["Success Tutoring - Belmont WA", 28, 0, 0, 4],
    ["Success Tutoring - Torquay", 10, 0, 0, 0],
])
INHOUSE_REVENUE = xlsx([
    ["Revenue"],
    ["Location Name", "Sessions", "Gross Revenue", "Net Revenue (ex tax)"],
    ["Burwood", 44, 4863.23, 4429.79],
    ["Belmont WA", 0, 103.01, 93.83],
    ["Landsdale", 0, 864.5, 785.91],
])

MASTER = pd.DataFrame([
    {"Location": "Success Tutoring - Auburn", "Country": "Australia"},
    {"Location": "Success Tutoring - Epping", "Country": "Australia"},
    {"Location": "Success Tutoring - Epping VIC", "Country": "Australia"},
    {"Location": "Success Tutoring - Belmont", "Country": "Australia"},
    {"Location": "Success Tutoring - Belmont WA", "Country": "Australia"},
    {"Location": "Success Tutoring - Howick", "Country": "New Zealand"},
    {"Location": "Success Tutoring - Burwood", "Country": "Australia"},
    {"Location": "Success Tutoring - Landsdale", "Country": "Australia"},
])


def all_exports():
    return [
        ingest.read_export("Weekly Membership.xlsx", HAPANA_MEMBERS),
        ingest.read_export("SA - Revenue Location Summary (16).xlsx", HAPANA_REV_AU),
        ingest.read_export("SA - Revenue Location Summary (17).xlsx", HAPANA_REV_NZ),
        ingest.read_export("members-report-2026-09-14-to-2026-09-20.xlsx", INHOUSE_MEMBERS),
        ingest.read_export("revenue-report-2026-09-14-to-2026-09-20.xlsx", INHOUSE_REVENUE),
    ]


def test_detects_each_export():
    kinds = [e["kind"] for e in all_exports()]
    assert kinds == [ingest.HAPANA_MEMBERS, ingest.HAPANA_REVENUE, ingest.HAPANA_REVENUE,
                     ingest.INHOUSE_MEMBERS, ingest.INHOUSE_REVENUE]
    assert all(e["week_end"] == date(2026, 9, 20) for e in all_exports())


def test_week_labels_match_sheet():
    # Pairs taken from the existing Weekly Membership / Revenue tabs.
    assert ingest.week_label(date(2022, 7, 10)) == "28/2022"
    assert ingest.week_label(date(2024, 1, 7)) == "1/2024"
    assert ingest.week_label(date(2026, 9, 20)) == "38/2026"
    assert ingest.week_label_to_date("38/2026") == date(2026, 9, 20)
    assert ingest.week_label_to_date("28/2022") == date(2022, 7, 10)


def test_combine_week_38():
    res = ingest.combine(all_exports(), MASTER)
    assert res["errors"] == []
    # Torquay isn't in the master list yet: flagged as a new location.
    assert list(res["unknown"]) == ["Success Tutoring - Torquay"]
    assert res["week_end"] == date(2026, 9, 20)

    mem = res["membership"].set_index("Location")
    assert "Success Tutoring - Success Tutoring Sandbox" not in mem.index
    # Epping / Epping VIC and Belmont / Belmont WA stay separate.
    assert mem.loc["Success Tutoring - Belmont", "active"] == 17
    assert mem.loc["Success Tutoring - Belmont WA", "active"] == 28
    assert mem.loc["Success Tutoring - Epping VIC", "active"] == 46

    rev = res["revenue"].set_index("Location")
    assert rev.loc["Success Tutoring - Auburn", "gross"] == 5970
    assert rev.loc["Success Tutoring - Auburn", "net"] == round(5970 / 1.1, 2)
    assert rev.loc["Success Tutoring - Howick", "net"] == round(1346 / 1.15, 2)
    # In-house net is used as given; members come from the members report.
    assert rev.loc["Success Tutoring - Burwood", "net"] == 4429.79
    assert rev.loc["Success Tutoring - Burwood", "active"] == 56
    assert rev.loc["Success Tutoring - Burwood", "visits"] == round(56 * 1.96, 1)
    assert rev.loc["Success Tutoring - Landsdale", "active"] == 0
    assert any("Landsdale" in i for i in res["info"])
    assert any("Sandbox" in i for i in res["info"])


def test_unknown_location_gets_suggestion():
    master = MASTER[MASTER["Location"] != "Success Tutoring - Belmont WA"]
    res = ingest.combine(all_exports(), master)
    assert "Success Tutoring - Belmont WA" in res["unknown"]
    assert "Belmont WA" in res["unknown"]
    assert "Success Tutoring - Belmont" in res["unknown"]["Belmont WA"]["suggestions"]


def test_alias_maps_other_spelling():
    exports = [ingest.read_export("members-report-2026-09-14-to-2026-09-20.xlsx", xlsx([
        ["Location Name", "Active Members", "Suspended Members", "Cancelled Members",
         "New Members"],
        ["ST Burwood", 56, 11, 3, 1],
    ]))]
    res = ingest.combine(exports, MASTER, {"ST Burwood": "Success Tutoring - Burwood"})
    assert res["unknown"] == {}
    assert res["membership"]["Location"].tolist() == ["Success Tutoring - Burwood"]


def test_two_names_on_one_site_is_an_error():
    res = ingest.combine(all_exports(), MASTER,
                         {"Belmont WA": "Success Tutoring - Belmont"})
    assert any("Belmont" in e and "more than once" in e for e in res["errors"])


def test_missing_gst_rule_is_an_error():
    master = MASTER.copy()
    master.loc[master["Location"] == "Success Tutoring - Howick", "Country"] = ""
    res = ingest.combine(all_exports(), master)
    assert any("Howick" in e and "GST" in e for e in res["errors"])


def test_rejects_monthly_and_unknown_files():
    with pytest.raises(ingest.ExportError, match="longer than one week"):
        ingest.read_export("revenue-report-2026-08-01-to-2026-08-31.xlsx", INHOUSE_REVENUE)
    with pytest.raises(ingest.ExportError, match="not recognised"):
        ingest.read_export("other.xlsx", xlsx([["Full Name", "Email"], ["x", "y"]]))


def test_hapana_export_with_data_on_second_tab():
    wb = openpyxl.Workbook()
    info = wb.active
    info.title = "Export info"
    for r in [["Exported at", "28/09/2026, 5:37 am (GMT+00:00)"], ["Filters"],
              ["Date range", "21/09/2026 - 27/09/2026"]]:
        info.append(r)
    data = wb.create_sheet("Weekly Membership")
    for r in [["Success Tutoring - Business name", "Date - Week/Year", "# Active members",
               "# Suspended members", "# Cancelled members", "# New members"],
              ["Success Tutoring - Auburn", "39/2026", 84, 9, 2, 3]]:
        data.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    e = ingest.read_export("Weekly Membership (4).xlsx", buf.getvalue())
    assert e["kind"] == ingest.HAPANA_MEMBERS
    assert e["week_end"] == date(2026, 9, 27)
    assert e["rows"].iloc[0]["active"] == 84
    assert e["warnings"] == []


def test_csv_export():
    data = (b"Location Name,Sessions,Gross Revenue,Net Revenue (ex tax)\n"
            b"Burwood,44,\"4,863.23\",$4429.79\n")
    e = ingest.read_export("revenue-report-2026-09-14-to-2026-09-20.csv", data)
    assert e["rows"].iloc[0]["gross"] == 4863.23
    assert e["rows"].iloc[0]["net"] == 4429.79


def test_sheet_records():
    res = ingest.combine(all_exports(), MASTER)
    recs = ingest.revenue_records(res["revenue"], res["week_end"])
    r = next(x for x in recs if x["Location"] == "Success Tutoring - Auburn")
    assert r["Date - Week/Year"] == "38/2026" and r["Date"] == "20/9/2026"
    assert r["Revenue per Session"] == round(round(5970 / 1.1, 2) / 36, 2)


def test_nz_gst_candidates():
    df = pd.DataFrame([
        {"Country": "New Zealand", "Gross Revenue": "$1,346", "Net Revenue": "$1,224"},
        {"Country": "New Zealand", "Gross Revenue": "$1,346", "Net Revenue": "$1,170"},
        {"Country": "New Zealand", "Gross Revenue": "$2,379", "Net Revenue": "$2,074"},
        {"Country": "Australia", "Gross Revenue": "$297", "Net Revenue": "$270"},
    ])
    out = ingest.nz_gst_candidates(df)
    assert out.index.tolist() == [0]
    assert out.iloc[0]["New Net"] == round(1346 / 1.15, 2)


def test_duplicate_weeks():
    df = pd.DataFrame({"Location": ["A", "A", "B"], "Date - Week/Year": ["1/2024"] * 3})
    out = ingest.duplicate_weeks(df, "Location")
    assert out["Location"].tolist() == ["A"]


def test_hapana_region_separates_same_names():
    # Current Hapana layout: Country/Region filled on the first row of each
    # group, both Eppings called plain "Epping".
    data = xlsx([
        ["Country", "Region", "Business name", "Date - Week/Year", "# Active members",
         "# Suspended members", "# Cancelled members", "# New members"],
        ["Australia", "New South Wales", "Auburn", "38/2026", 83, 9, 3, 1],
        [None, None, "Epping", "38/2026", 100, 12, 0, 2],
        [None, "Victoria", "Belmont", "38/2026", 17, 3, 0, 0],
        [None, None, "Epping", "38/2026", 46, 10, 0, 1],
        ["New Zealand", "New Zealand", "Howick", "38/2026", 19, 3, 0, 0],
    ])
    master = MASTER.assign(Region=["New South Wales", "New South Wales", "Victoria",
                                   "Victoria", "Western Australia", "Auckland",
                                   "New South Wales", "Western Australia"])
    res = ingest.combine([ingest.read_export("Weekly Membership (4).xlsx", data)], master)
    assert res["errors"] == [] and res["unknown"] == {}
    mem = res["membership"].set_index("Location")["active"]
    assert mem["Success Tutoring - Epping"] == 100
    assert mem["Success Tutoring - Epping VIC"] == 46
    assert mem["Success Tutoring - Belmont"] == 17
    assert mem["Success Tutoring - Howick"] == 19


def test_region_mismatch_is_flagged_not_merged():
    data = xlsx([
        ["Country", "Region", "Business name", "Date - Week/Year", "# Active members",
         "# Suspended members", "# Cancelled members", "# New members"],
        ["Australia", "Queensland", "Epping", "38/2026", 5, 0, 0, 0],
    ])
    master = MASTER.assign(Region="New South Wales")
    res = ingest.combine([ingest.read_export("wm.xlsx", data)], master)
    assert list(res["unknown"]) == ["Epping QLD"]


def test_inhouse_dash_names_and_total_rows():
    data = xlsx([
        ["Location Name", "Sessions", "Gross Revenue", "Net Revenue (ex tax)"],
        ["Success Tutoring – Burwood", 44, 4863.23, 4429.79],
        ["Total (AUD) — 1 centres", 44, 4863.23, 4429.79],
        ["Total (NZD) — 3 centres", 0, 0, 0],
    ])
    res = ingest.combine([ingest.read_export("revenue-report-2026-09-14-to-2026-09-20.xlsx",
                                             data)], MASTER)
    assert res["unknown"] == {}
    assert res["revenue"]["Location"].tolist() == ["Success Tutoring - Burwood"]


RAW_INHOUSE_MEMBERS = xlsx([
    ["Location Name", "Active Members", "Suspended Members", "Cancelled Members",
     "New Members"],
    ["Success Tutoring - Belmont", 28, 0, 0, 4],
    ["Success Tutoring – Burwood", 56, 11, 3, 1],
])


def test_same_name_in_both_crms_is_a_clash_to_resolve():
    exports = [ingest.read_export("members-report-2026-09-14-to-2026-09-20.xlsx",
                                  RAW_INHOUSE_MEMBERS),
               ingest.read_export("Weekly Membership.xlsx", HAPANA_MEMBERS)]
    res = ingest.combine(exports, MASTER)
    assert [(c["crm"], c["name"], c["location"]) for c in res["clashes"]] == [
        ("In-house", "Success Tutoring - Belmont", "Success Tutoring - Belmont")]
    assert "Success Tutoring - Belmont WA" in res["clashes"][0]["suggestions"]

    # Mapping for the in-house CRM only: Hapana's Belmont is unaffected.
    alias = {ingest.scoped_alias("In-house", "Success Tutoring - Belmont"):
             "Success Tutoring - Belmont WA"}
    res = ingest.combine(exports, MASTER, alias)
    assert res["errors"] == [] and res["clashes"] == []
    mem = res["membership"].set_index("Location")["active"]
    assert mem["Success Tutoring - Belmont WA"] == 28
    assert mem["Success Tutoring - Belmont"] == 17


def test_scoped_alias_keys():
    assert ingest.alias_key("In-house: Success Tutoring - Belmont") == "in-house:belmont"
    assert ingest.alias_key("Belmont WA") == "belmont wa"
    assert ingest.canonical_new_name("In-house: Epping") == "Success Tutoring - Epping"


def test_region_spelling_differences_still_match():
    data = xlsx([
        ["Country", "Region", "Business name", "Date - Week/Year", "# Active members",
         "# Suspended members", "# Cancelled members", "# New members"],
        ["Australia", "Australian Capital Territory", "Dickson", "38/2026", 47, 10, 3, 1],
        [None, "Victoria", "Epping", "38/2026", 46, 10, 0, 1],
    ])
    master = pd.DataFrame([
        {"Location": "Success Tutoring - Dickson", "Region": "ACT", "Country": "Australia"},
        {"Location": "Success Tutoring - Epping", "Region": "NSW", "Country": "Australia"},
        {"Location": "Success Tutoring - Epping VIC", "Region": "", "Country": "Australia"},
    ])
    res = ingest.combine([ingest.read_export("wm.xlsx", data)], master)
    assert res["unknown"] == {}
    assert res["membership"]["Location"].tolist() == [
        "Success Tutoring - Dickson", "Success Tutoring - Epping VIC"]


def test_empty_inhouse_rows_are_left_out():
    members = xlsx([
        ["Location Name", "Active Members", "Suspended Members", "Cancelled Members",
         "New Members"],
        ["Success Tutoring - Epping", 0, 0, 0, 0],
        ["Success Tutoring - Head Office", 0, 0, 0, 0],
        ["Success Tutoring - Burwood", 56, 11, 3, 1],
        ["Success Tutoring - Landsdale", 0, 0, 0, 0],
    ])
    revenue = xlsx([
        ["Location Name", "Sessions", "Gross Revenue", "Net Revenue (ex tax)"],
        ["Success Tutoring - Epping", 0, 0, 0],
        ["Success Tutoring - Burwood", 44, 4863.23, 4429.79],
        ["Success Tutoring - Landsdale", 0, 864.5, 785.91],   # presale: revenue only
    ])
    master = MASTER.assign(Status="Trading")
    exports = [ingest.read_export("members-report-2026-09-14-to-2026-09-20.xlsx", members),
               ingest.read_export("revenue-report-2026-09-14-to-2026-09-20.xlsx", revenue),
               ingest.read_export("Weekly Membership.xlsx", HAPANA_MEMBERS)]
    res = ingest.combine(exports, master)
    assert res["errors"] == [] and res["clashes"] == [] and res["unknown"] == {}
    mem = res["membership"].set_index("Location")["active"]
    assert mem["Success Tutoring - Epping"] == 100          # Hapana's Epping only
    assert "Success Tutoring - Landsdale" in mem.index      # kept: has revenue
    assert any("Head Office" in i and "Epping" in i for i in res["info"])
    assert not any("Epping" in w for w in res["warnings"])  # Epping has Hapana data


def test_trading_site_with_nothing_is_warned():
    members = xlsx([
        ["Location Name", "Active Members", "Suspended Members", "Cancelled Members",
         "New Members"],
        ["Success Tutoring - Burwood", 0, 0, 0, 0],
    ])
    res = ingest.combine([ingest.read_export(
        "members-report-2026-09-14-to-2026-09-20.xlsx", members)],
        MASTER.assign(Status="Trading"))
    assert any("Burwood is Trading" in w for w in res["warnings"])


def test_hapana_revenue_row_shared_by_two_sites_is_left_out():
    members = xlsx([
        ["Country", "Region", "Business name", "Date - Week/Year", "# Active members",
         "# Suspended members", "# Cancelled members", "# New members"],
        ["Australia", "New South Wales", "Epping", "38/2026", 100, 12, 0, 2],
        [None, "Victoria", "Epping", "38/2026", 46, 10, 0, 1],
        [None, None, "Belmont", "38/2026", 17, 3, 0, 0],
    ])
    revenue = xlsx([
        ["Business name", "Date - Week/Year", "# Active members - Date range end",
         "Total Sessions", "Average Gross Revenue per Location"],
        ["Epping", "38/2026", 146, 45, "10,685"],
        ["Belmont", "38/2026", 17, 13, "1,088"],
    ])
    master = MASTER.assign(Region=["New South Wales", "New South Wales", "Victoria",
                                   "Victoria", "Western Australia", "Auckland",
                                   "New South Wales", "Western Australia"])
    res = ingest.combine([ingest.read_export("Weekly Membership.xlsx", members),
                          ingest.read_export("SA - Revenue Location Summary.xlsx", revenue)],
                         master)
    assert res["errors"] == []
    assert res["revenue"]["Location"].tolist() == ["Success Tutoring - Belmont"]
    w = next(w for w in res["warnings"] if "one 'Epping' row" in w)
    assert "Success Tutoring - Epping, Success Tutoring - Epping VIC" in w


def test_hapana_revenue_with_region_and_no_member_column():
    # Revenue report after adding Region and removing the members column.
    members = xlsx([
        ["Country", "Region", "Business name", "Date - Week/Year", "# Active members",
         "# Suspended members", "# Cancelled members", "# New members"],
        ["Australia", "New South Wales", "Epping", "37/2026", 100, 12, 0, 2],
        [None, "Victoria", "Belmont", "37/2026", 17, 3, 0, 0],
        [None, None, "Epping", "37/2026", 46, 10, 0, 1],
        ["New Zealand", "New Zealand", "Howick", "37/2026", 19, 3, 0, 0],
    ])
    revenue = xlsx([
        ["SA - Revenue Location Summary"],
        ["Business name", "Region", "Date - Week/Year", "Total Sessions",
         "Average Gross Revenue per Location"],
        ["Belmont", "Victoria", "37/2026", 12, "1,334"],
        ["Epping", "New South Wales", "37/2026", 23, "7,152"],
        ["Epping", "Victoria", "37/2026", 22, "3,533"],
        ["Howick", "New Zealand", "37/2026", 17, "1,346"],
    ])
    master = MASTER.assign(Region=["New South Wales", "New South Wales", "Victoria",
                                   "Victoria", "Western Australia", "Auckland",
                                   "New South Wales", "Western Australia"])
    res = ingest.combine([ingest.read_export("Weekly Membership.xlsx", members),
                          ingest.read_export("SA - Revenue Location Summary (19).xlsx", revenue)],
                         master)
    assert res["errors"] == [] and res["unknown"] == {}
    assert not any("one 'Epping' row" in w for w in res["warnings"])
    rev = res["revenue"].set_index("Location")
    assert rev.loc["Success Tutoring - Epping", "gross"] == 7152
    assert rev.loc["Success Tutoring - Epping VIC", "gross"] == 3533
    assert rev.loc["Success Tutoring - Epping VIC", "active"] == 46     # from members
    assert rev.loc["Success Tutoring - Epping VIC", "net"] == round(3533 / 1.1, 2)
    assert rev.loc["Success Tutoring - Howick", "net"] == round(1346 / 1.15, 2)
    assert rev.loc["Success Tutoring - Howick", "active"] == 19


def test_hapana_blank_repeated_business_name():
    revenue = xlsx([
        ["Business name", "Region", "Date - Week/Year", "Total Sessions",
         "Average Gross Revenue per Location"],
        ["Ellenbrook", "Western Australia", "38/2026", 108, "20,350"],
        ["Epping", "New South Wales", "38/2026", 23, "7,079"],
        [None, "Victoria", "38/2026", 22, "3,545"],
        ["Forrestfield", "Western Australia", "38/2026", 43, "4,671"],
        [None, None, None, None, None],
    ])
    e = ingest.read_export("SA - Revenue Location Summary (20).xlsx", revenue)
    assert [(r["name"], r["region"], r["gross"]) for _, r in e["rows"].iterrows()] == [
        ("Ellenbrook", "Western Australia", 20350), ("Epping", "New South Wales", 7079),
        ("Epping", "Victoria", 3545), ("Forrestfield", "Western Australia", 4671)]
    master = MASTER.assign(Region=["New South Wales", "New South Wales", "Victoria",
                                   "Victoria", "Western Australia", "Auckland",
                                   "New South Wales", "Western Australia"])
    rev = ingest.combine([e], master)["revenue"].set_index("Location")["gross"]
    assert rev["Success Tutoring - Epping VIC"] == 3545
    assert rev["Success Tutoring - Epping"] == 7079
