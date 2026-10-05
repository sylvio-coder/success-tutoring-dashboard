"""Monthly P&L workbook: reading, checks, and the P&L upload page."""
import io

import pandas as pd
import pytest
from openpyxl import Workbook

import pnl
from fake_sheet import FakeSpreadsheet
from test_admin_pages import REV, VLOOKUP, WM, Upload, run_page

ACCOUNTS = [("", "Revenue"), (200, "Sales"), (260, "Other Revenue"), (270, "Interest Income"),
            ("", "Total Revenue"), (None, None), ("", "Cost of Sales"),
            (472, "Wages - Tutor"), (474, "Superannuation - Tutor"),
            ("", "Operating Costs"), (469, "Rent"), (467, "Royalties"), (400, "Marketing"),
            (453, "Office Expenses"), ("", "EBITDA"),
            ("", "Interest, Tax & Depreciation"), (437, "Interest Expense"),
            ("", "Memo — not included in Net Profit"), (479, "Unrelated Expenses")]
CENTRES = {  # name (currency): amounts by code
    "Auburn (AUD)": {200: 20000, 270: 33, 472: 6000, 474: 700, 469: 4000, 467: 2200,
                     400: 900, 453: 100, 437: 500, 479: 999},
    "Burwod (AUD)": {200: 900, 469: 3000, 453: -50},
    "Howick (NZD)": {200: 8000, 472: 2500, 469: 2000, 467: 2200, 400: 300},
}


def workbook(title="P&L by centre — August 2026 — each centre in its own currency",
             centres=CENTRES, summary=True):
    wb = Workbook()
    s = wb.active
    s.title = "Summary"
    if summary:
        s.append(["Network P&L — August 2026"])
        s.append([])
        s.append([])
        s.append(["Centre", "Country", "Currency", "Status", "Revenue", "Expenses",
                  "Net profit", "Margin", "FX rate to AUD"])
        s.append(["Auburn", "Australia", "AUD", "Submitted", 0, 0, 0, 0, 1])
        s.append(["Burwod", "Australia", "AUD", "Amended", 0, 0, 0, 0, 1])
        s.append(["Howick", "New Zealand", "NZD", "Submitted", 0, 0, 0, 0, 0.94])
        s.append(["Ryde", "Australia", "AUD", "Pending", None, None, None, None, None])
        s.append(["Network total (AUD)", "", "", "", None])
    ws = wb.create_sheet("P&L (native currency)")
    ws.append([title])
    ws.append(["3 declared centres."])
    ws.append([])
    ws.append(["Code", "Account"] + list(centres))
    for code, name in ACCOUNTS:
        if code is None:
            ws.append([])
        elif code == "":
            ws.append(["", name] + [None] * len(centres))
        else:
            ws.append([code, name] + [c.get(code, 0) for c in centres.values()])
    wb.create_sheet("P&L (AUD)").append(["Code", "Account", "Auburn"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def master():
    return pd.DataFrame([dict(zip(VLOOKUP[0], r)) for r in VLOOKUP[1:]])


def test_reads_native_tab_with_summary():
    d = pnl.read_pnl("success-insights-pnl-2026-08.xlsx", workbook())
    assert d["period"] == "2026-08" and d["source"] == "P&L (native currency)"
    assert d["pending"] == ["Ryde"]
    c = d["centres"].set_index("Centre")
    assert c.loc["Howick", "Currency"] == "NZD" and c.loc["Howick", "FX rate to AUD"] == 0.94
    assert c.loc["Burwod", "Status"] == "Amended"
    a = c.loc["Auburn"]
    assert a["Revenue"] == 20033
    assert a["Expenses"] == 6000 + 700 + 4000 + 2200 + 900 + 100   # no interest, no 479
    assert a["EBITDA"] == 20033 - 13900 and a["Net Profit"] == a["EBITDA"] - 500
    sec = d["lines"].set_index(["Centre", "Account code"])["P&L section"]
    assert sec[("Auburn", 472)] == "Cost of Sales" and sec[("Auburn", 469)] == "Operating Costs"
    assert sec[("Auburn", 479)].startswith("Memo")


def test_period_from_file_name_and_no_summary():
    d = pnl.read_pnl("pnl-2026-09.xlsx", workbook(title="P&L by centre", summary=False))
    assert d["period"] == "2026-09"
    assert d["centres"].set_index("Centre").loc["Howick", "Currency"] == "NZD"


def test_not_a_pnl_file():
    wb = Workbook()
    wb.active.append(["Hello"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(pnl.PnlError):
        pnl.read_pnl("x.xlsx", buf.getvalue())


def test_matching_checks_and_records():
    d = pnl.read_pnl("p.xlsx", workbook())
    c = pnl.match_centres(d["centres"], master(), {})
    loc = c.set_index("Centre")["Location"]
    assert loc["Auburn"] == "Success Tutoring - Auburn"
    assert pd.isna(loc["Burwod"])          # the P&L spells it differently
    assert "Success Tutoring - Burwood" in c.set_index("Centre").loc["Burwod", "suggestions"]
    c = pnl.match_centres(d["centres"], master(),
                          {"P&L: Burwod": "Success Tutoring - Burwood"})
    found = pnl.checks(d["lines"], c, master())
    by = found.groupby("Location")["Check"].apply(set)
    assert by["Success Tutoring - Burwood"] >= {"Very low revenue", "No royalties",
                                                   "Negative costs (credits)"}
    assert "Success Tutoring - Auburn" not in by.index
    assert "Success Tutoring - Howick" not in by.index     # NZ in both, has a rate to AUD
    m = master()
    m.loc[m["Location"] == "Success Tutoring - Howick", "Country"] = "Australia"
    assert "Country differs" in set(pnl.checks(d["lines"], c, m)["Check"])
    recs = pnl.records(d["lines"], c, "2026-08")
    assert {r["Location"] for r in recs} == {"Success Tutoring - Auburn",
                                             "Success Tutoring - Burwood",
                                             "Success Tutoring - Howick"}
    r = next(r for r in recs if r["Location"].endswith("Howick") and r["Account code"] == 200)
    assert r["Amount"] == 8000 and r["Currency"] == "NZD" and r["FX rate to AUD"] == 0.94


def test_aliases_for_pnl_do_not_affect_crm_uploads():
    import ingest
    m = ingest.LocationMatcher(["Success Tutoring - Belmont", "Success Tutoring - Belmont WA"],
                               {"P&L: Belmont": "Success Tutoring - Belmont WA"})
    assert m.resolve("Belmont", crm="P&L") == "Success Tutoring - Belmont WA"
    assert m.resolve("Belmont", crm="Hapana") == "Success Tutoring - Belmont"


@pytest.fixture
def ss():
    return FakeSpreadsheet({"Vlookup": VLOOKUP, "Weekly Membership": WM, "Revenue": REV})


def test_upload_page_maps_name_and_writes_month(ss):
    files = Upload("success-insights-pnl-2026-08.xlsx", workbook())
    at = run_page("page_pnl_upload", ss, files)
    assert not at.exception, at.exception
    confirm = next(b for b in at.button if b.label.startswith("✅ Confirm"))
    assert "August 2026" in confirm.label and confirm.disabled      # Burwod unknown
    at.selectbox(key="pnl_map_Burwod").set_value("Success Tutoring - Burwood").run()
    confirm = next(b for b in at.button if b.label.startswith("✅ Confirm"))
    assert not confirm.disabled
    confirm.click().run()
    assert not at.exception, at.exception
    assert any("August 2026 written" in s.value for s in at.success)

    tab = ss.worksheet("P&L").cells
    assert tab[0] == pnl.PNL_HEADERS
    assert len(tab) - 1 == sum(1 for code, _ in ACCOUNTS if isinstance(code, int)) * 3
    assert {r[0] for r in tab[1:]} == {"2026-08"}
    assert "Success Tutoring - Burwood" in {r[1] for r in tab[1:]}
    assert ss.worksheet("Location Aliases").cells[-1] == ["P&L: Burwod",
                                                          "Success Tutoring - Burwood"]
    assert ss.worksheet("Vlookup").cells[0][-1] == "Royalty Min"

    # The same month again replaces rather than duplicates, and the alias is remembered.
    at2 = run_page("page_pnl_upload", ss, files)
    assert any("will be replaced" in i.value for i in at2.info)
    next(b for b in at2.button if b.label.startswith("✅ Confirm")).click().run()
    assert not at2.exception, at2.exception
    assert len(ss.worksheet("P&L").cells) == len(tab)


def test_reupload_warns_about_centres_it_wont_change(ss):
    run_page("page_pnl_upload", ss, Upload("p.xlsx", workbook()))
    at = run_page("page_pnl_upload", ss, Upload("p.xlsx", workbook()))
    at.selectbox(key="pnl_map_Burwod").set_value("Success Tutoring - Burwood").run()
    next(b for b in at.button if b.label.startswith("✅ Confirm")).click().run()
    fewer = {k: v for k, v in CENTRES.items() if not k.startswith("Howick")}
    at = run_page("page_pnl_upload", ss, Upload("p.xlsx", workbook(centres=fewer)))
    assert not at.exception, at.exception
    assert any("won't change: Howick" in w.value for w in at.warning)


def test_upload_page_leave_out(ss):
    at = run_page("page_pnl_upload", ss, Upload("p.xlsx", workbook()))
    at.selectbox(key="pnl_map_Burwod").set_value("Leave out this month").run()
    next(b for b in at.button if b.label.startswith("✅ Confirm")).click().run()
    assert not at.exception, at.exception
    assert "Success Tutoring - Burwood" not in {r[1] for r in ss.worksheet("P&L").cells[1:]}


def test_locations_page_shows_royalty_min():
    vl = [VLOOKUP[0] + ["Royalty Min"]] + [r + [""] for r in VLOOKUP[1:]]
    vl[1][-1] = "$2,266"
    ss = FakeSpreadsheet({"Vlookup": vl, "Weekly Membership": WM, "Revenue": REV})
    at = run_page("page_locations", ss)
    assert not at.exception, at.exception
    editor = at.dataframe[0].value
    assert editor.set_index("Location").loc["Success Tutoring - Auburn", "Royalty Min"] == 2266
    assert not any("unsaved change" in i.value for i in at.info)    # parsing isn't an edit
